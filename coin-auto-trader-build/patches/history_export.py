"""
투자 이력 정리 · 내보내기.

DB 에 쌓인 기록을 사람이 읽는 표로 만들고, 엑셀(.xlsx) 또는 CSV 로 저장한다.

- openpyxl 이 있으면 진짜 엑셀 파일(.xlsx)로, 시트를 나눠서 저장한다.
- 없으면 CSV 로 저장한다. 이때 UTF-8 BOM 을 붙인다 - BOM 이 없으면 엑셀이
  한글을 깨서 연다.
"""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None


TRADE_HEADERS = ["번호", "시각", "종목", "구분", "체결가(원)", "수량",
                 "거래금액(원)", "수수료(원)", "실현손익(원)", "모드", "전략", "매매 사유"]
SESSION_HEADERS = ["번호", "시작 시각", "종료 시각", "구동 시간", "모드",
                   "배정 예산(원)", "시작 자산(원)", "종료 자산(원)",
                   "구간 손익(원)", "매매 횟수", "종료 사유"]
RISE_HEADERS = ["시각", "종목", "이전가(원)", "현재가(원)", "상승률(%)",
                "관측 구간(분)", "보유 중이었나"]
DAILY_HEADERS = ["날짜", "매매 횟수", "매수", "매도", "실현손익(원)", "승률(%)"]


def _fmt_dt(iso: Optional[str]) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso).strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        return str(iso)


def _duration_text(start_iso: str, end_iso: Optional[str]) -> str:
    """'2시간 13분' 형태의 구동 시간."""
    if not start_iso:
        return ""
    try:
        start = datetime.fromisoformat(start_iso)
        end = datetime.fromisoformat(end_iso) if end_iso else datetime.now(start.tzinfo)
    except (ValueError, TypeError):
        return ""
    return humanize_seconds((end - start).total_seconds())


def humanize_seconds(seconds: float) -> str:
    seconds = max(int(seconds), 0)
    days, rem = divmod(seconds, 86400)
    hours, rem = divmod(rem, 3600)
    minutes, secs = divmod(rem, 60)
    parts = []
    if days:
        parts.append(f"{days}일")
    if hours or days:
        parts.append(f"{hours}시간")
    if minutes or hours or days:
        parts.append(f"{minutes}분")
    if not parts:
        parts.append(f"{secs}초")
    return " ".join(parts)


def session_seconds(start_iso: str, end_iso: Optional[str]) -> float:
    try:
        start = datetime.fromisoformat(start_iso)
        end = datetime.fromisoformat(end_iso) if end_iso else datetime.now(start.tzinfo)
        return max((end - start).total_seconds(), 0.0)
    except (ValueError, TypeError):
        return 0.0


# --------------------------------------------------------------------------
# 표 만들기
# --------------------------------------------------------------------------

def trade_rows(db) -> list[list]:
    rows = []
    for i, t in enumerate(db.all_trades(), 1):
        rows.append([
            i, _fmt_dt(t["ts"]), t["market"],
            "매수" if t["side"] == "buy" else "매도",
            round(t["price"], 2), round(t["volume"], 8),
            round(t["krw_amount"], 0), round(t["fee_krw"], 0),
            round(t["realized_pnl_krw"], 0) if t["realized_pnl_krw"] is not None else "",
            "실거래" if t["mode"] == "live" else "모의",
            t["strategy"] or "", t["reason"] or "",
        ])
    return rows


def session_rows(db) -> list[list]:
    rows = []
    for i, s in enumerate(db.all_sessions(), 1):
        rows.append([
            i, _fmt_dt(s["started_at"]),
            _fmt_dt(s["ended_at"]) if s["ended_at"] else "(구동 중)",
            _duration_text(s["started_at"], s["ended_at"]),
            "실거래" if s["mode"] == "live" else "모의",
            round(s["budget_krw"] or 0),
            round(s["start_value_krw"] or 0),
            round(s["end_value_krw"] or 0) if s["end_value_krw"] is not None else "",
            round(s["realized_pnl_krw"] or 0),
            s["num_trades"] or 0, s["stop_reason"] or "",
        ])
    return rows


def rise_rows(db, limit: int = 500) -> list[list]:
    rows = []
    for r in db.rise_events(limit=limit):
        rows.append([
            _fmt_dt(r["ts"]), r["market"], round(r["price_from"], 2),
            round(r["price_to"], 2), round(r["change_pct"], 2),
            r["window_min"], "예" if r["held"] else "아니오",
        ])
    return rows


def daily_rows(db) -> list[list]:
    """날짜별 매매 요약."""
    buckets: dict[str, dict] = {}
    for t in db.all_trades():
        day = (_fmt_dt(t["ts"]) or "")[:10]
        if not day:
            continue
        b = buckets.setdefault(day, {"n": 0, "buy": 0, "sell": 0, "pnl": 0.0, "win": 0})
        b["n"] += 1
        if t["side"] == "buy":
            b["buy"] += 1
        else:
            b["sell"] += 1
            pnl = t["realized_pnl_krw"] or 0
            b["pnl"] += pnl
            if pnl > 0:
                b["win"] += 1
    rows = []
    for day in sorted(buckets):
        b = buckets[day]
        win_rate = round(b["win"] / b["sell"] * 100, 1) if b["sell"] else ""
        rows.append([day, b["n"], b["buy"], b["sell"], round(b["pnl"]), win_rate])
    return rows


def daily_pnl_series(db) -> list[tuple]:
    """막대그래프용 (MM/DD, 실현손익) 목록."""
    return [(r[0][5:].replace("-", "/"), float(r[4])) for r in daily_rows(db)]


def equity_series(db, limit: int = 300) -> list[tuple]:
    """자산 추이 선그래프용 (시각, 평가액) 목록."""
    snaps = db.ledger_snapshots_since("")
    if len(snaps) > limit:            # 너무 많으면 균등하게 솎아낸다
        stride = -(-len(snaps) // limit)
        snaps = snaps[::stride]
    out = []
    for s in snaps:
        try:
            label = datetime.fromisoformat(s["ts"]).strftime("%m/%d %H:%M")
        except (ValueError, TypeError):
            label = str(s["ts"])[:16]
        out.append((label, float(s["total_value_krw"])))
    return out


# --------------------------------------------------------------------------
# 파일로 내보내기
# --------------------------------------------------------------------------

def _autosize(ws) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(width + 4, 10), 52)


def _write_sheet(wb, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title=title)
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="DDE6F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    _autosize(ws)


def export(db, out_dir: Path, filename_stem: str = "") -> Path:
    """이력 전체를 파일로 저장하고 저장 경로를 돌려준다."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = filename_stem or ("투자이력_" + datetime.now().strftime("%Y%m%d_%H%M"))

    sheets = [
        ("매매 이력", TRADE_HEADERS, trade_rows(db)),
        ("일별 요약", DAILY_HEADERS, daily_rows(db)),
        ("구동 이력", SESSION_HEADERS, session_rows(db)),
        ("상승 기록", RISE_HEADERS, rise_rows(db)),
    ]

    if openpyxl is not None:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)          # 기본 빈 시트 제거
        for title, headers, rows in sheets:
            _write_sheet(wb, title, headers, rows)
        path = out_dir / (stem + ".xlsx")
        wb.save(path)
        return path

    # openpyxl 이 없으면 CSV. 엑셀에서 한글이 깨지지 않도록 BOM 을 붙인다.
    path = out_dir / (stem + ".csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        for idx, (title, headers, rows) in enumerate(sheets):
            if idx:
                writer.writerow([])
            writer.writerow(["[" + title + "]"])
            writer.writerow(headers)
            writer.writerows(rows)
    return path
